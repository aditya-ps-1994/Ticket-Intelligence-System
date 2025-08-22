
import pandas as pd
from textblob import TextBlob
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, PieChart, Reference

# Load the Excel file
df = pd.read_excel("Ticket_dump_analysis.xlsx")

# Filter only English-language tickets
df_en = df[df['language'] == 'en'].copy()

# Apply sentiment analysis on the 'body' column
df_en['sentiment'] = df_en['body'].apply(lambda x: TextBlob(str(x)).sentiment.polarity)

# Save sentiment-enhanced dataset
df_en.to_excel("Ticket_dump_with_sentiment.xlsx", index=False)

# Prepare data for charts
type_counts = df_en['type'].value_counts().reset_index()
type_counts.columns = ['Type', 'Count']

queue_counts = df_en['queue'].value_counts().reset_index()
queue_counts.columns = ['Queue', 'Count']

priority_counts = df_en['priority'].value_counts().reset_index()
priority_counts.columns = ['Priority', 'Count']

# Create Excel workbook with charts
wb = Workbook()
ws_data = wb.active
ws_data.title = "Sentiment Data"

# Add full data
for r in dataframe_to_rows(df_en, index=False, header=True):
    ws_data.append(r)

# Add chart sheets
summary_data = [
    ("Ticket Type Count", type_counts, "Type", "Count", BarChart),
    ("Queue Count", queue_counts, "Queue", "Count", BarChart),
    ("Priority Breakdown", priority_counts, "Priority", "Count", PieChart)
]

for title, data, cat_col, val_col, ChartType in summary_data:
    ws = wb.create_sheet(title)
    for r in dataframe_to_rows(data, index=False, header=True):
        ws.append(r)
    chart = ChartType()
    chart.title = title
    data_ref = Reference(ws, min_col=2, min_row=1, max_row=len(data)+1)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(data)+1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "E2")

# Save final Excel output
wb.save("Ticket_Intelligence_Final.xlsx")
